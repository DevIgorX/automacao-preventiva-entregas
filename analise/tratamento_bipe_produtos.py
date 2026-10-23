"""Leitura do Bipe de produtos pronto ou bruto, sem modificar a planilha original."""
import json
import math
from numbers import Integral, Real
import re

import pandas as pd
from openpyxl.utils.cell import get_column_letter

STATUS_BIPADO = 'Bipado no depósito'


class BipeInvalido(ValueError):
    def __init__(self, mensagem, reconhecidos=0):
        super().__init__(mensagem)
        self.reconhecidos = reconhecidos


def _vazio(valor):
    return (valor is None or valor is pd.NA or valor is pd.NaT
        or isinstance(valor, Real) and not isinstance(valor, Integral) and math.isnan(valor)
        or isinstance(valor, str) and not valor.strip())


def normalizar_pedido_bipe(valor):
    """Mantém o identificador textual; Full com mais de 9 dígitos usa sufixo -1."""
    if _vazio(valor):
        raise ValueError('pedido vazio')
    if isinstance(valor, bool):
        raise ValueError('valor não é um número de pedido')
    if isinstance(valor, (Integral, Real)):
        if (not isinstance(valor, Integral) and not math.isfinite(valor)) or int(valor) != valor or valor <= 0:
            raise ValueError('número de pedido inválido')
        texto = str(int(valor))
        if len(texto) > 15:
            raise ValueError('pedido numérico com mais de 15 dígitos; obtenha o código original como texto para evitar dígitos arredondados pelo Excel')
    elif isinstance(valor, str):
        texto = valor.strip()
    else:
        raise ValueError('formato de pedido não reconhecido')

    # Trata apenas a parte decimal final, sem remover trechos no meio do código.
    decimal = re.fullmatch(r'([0-9]+)\.0+', texto)
    if decimal:
        texto = decimal.group(1)
    if not re.fullmatch(r'[0-9]+(?:-[0-9]+)?', texto) or not any(c in '123456789' for c in texto.split('-')[0]):
        raise ValueError('pedido deve conter números e, se houver, o sufixo após o hífen')
    full = re.fullmatch(r'([0-9]{10,})-[0-9]+', texto)
    return full.group(1) + '-1' if full else texto


def _extrair_pedido(valor):
    if isinstance(valor, str):
        texto = valor.strip()
        # Somente JSON estruturado ou string JSON: números puros seguem como texto.
        if texto.startswith(('{', '[', '"')):
            try:
                valor = json.loads(texto)
            except (ValueError, RecursionError) as exc:
                raise ValueError('JSON incompleto ou inválido') from exc
    if isinstance(valor, dict):
        order = valor.get('order')
        if isinstance(order, dict) and not _vazio(order.get('id')):
            valor = order['id']
        elif not _vazio(valor.get('external_grouper_code')):
            valor = valor['external_grouper_code']
        else:
            raise ValueError('JSON sem order.id ou external_grouper_code')
        # Inteiros dentro do JSON são exatos; não foram armazenados como número no Excel.
        if isinstance(valor, Integral) and not isinstance(valor, bool):
            valor = str(valor)
    return normalizar_pedido_bipe(valor)


def _erro_registros(aba, falhas, reconhecidos):
    exemplos = '; '.join(f'{celula}: {motivo}' for celula, motivo in falhas[:4])
    raise BipeInvalido(
        f'Bipe de produtos — aba "{aba}": não foi possível ler {len(falhas)} registro(s). '
        f'Confira {exemplos}. Corrija o arquivo e envie novamente.',
        reconhecidos=reconhecidos,
    )


def _ler_aba(livro, aba):
    bruto = pd.read_excel(livro, sheet_name=aba, header=None, dtype=object, keep_default_na=False)
    bruto = bruto.dropna(how='all')
    if bruto.empty:
        raise BipeInvalido('Bipe de produtos: a aba está vazia. Envie um relatório com dados.')

    # Reconhece o formato tratado pelos títulos da primeira linha.
    titulos = [str(v).strip().lower().title() for v in bruto.iloc[0]]
    if 'Pedido' in titulos:
        if titulos.count('Pedido') != 1 or titulos.count('Status_Deposito') > 1:
            raise BipeInvalido('Bipe de produtos: há títulos Pedido ou Status_Deposito repetidos. Confira as colunas.')
        df = bruto.iloc[1:].copy()
        df.columns = titulos
        pedidos = []
        indices = []
        falhas = []
        coluna_pedido = get_column_letter(titulos.index('Pedido') + 1)
        for indice, row in df.iterrows():
            if all(_vazio(v) for v in row):
                continue
            try:
                pedidos.append(normalizar_pedido_bipe(row['Pedido']))
                indices.append(indice)
            except ValueError as exc:
                falhas.append((f'{coluna_pedido}{indice + 1}', str(exc)))
        if falhas:
            _erro_registros(aba, falhas, len(pedidos))
        if not pedidos:
            raise BipeInvalido('Bipe de produtos: a aba está vazia. Envie um relatório com dados.')
        df = df.loc[indices].reset_index(drop=True)
        df['Pedido'] = pd.Series(pedidos, dtype=object)
        # Não sobrescreve os status de uma planilha já tratada.
        if 'Status_Deposito' not in df:
            df['Status_Deposito'] = STATUS_BIPADO
        return df

    # No arquivo bruto, a primeira linha também é um bipe, não um cabeçalho.
    pedidos = []
    falhas = []
    for linha, row in bruto.iterrows():
        for coluna, valor in enumerate(row, start=1):
            if _vazio(valor):
                continue
            try:
                pedidos.append(_extrair_pedido(valor))
            except ValueError as exc:
                falhas.append((f'{get_column_letter(coluna)}{linha + 1}', str(exc)))
    if falhas:
        _erro_registros(aba, falhas, len(pedidos))
    if not pedidos:
        raise BipeInvalido('Bipe de produtos: não foram encontrados pedidos nesta aba.')
    # Preserva a ordem e os bipes repetidos por produto/volume.
    return pd.DataFrame({'Pedido': pedidos, 'Status_Deposito': STATUS_BIPADO})


def ler_bipe_produtos(caminho, sheet_name=0):
    """Retorna as colunas para o cruzamento, aceitando JSON e planilhas tratadas."""
    with pd.ExcelFile(caminho) as livro:
        aba = livro.sheet_names[sheet_name] if isinstance(sheet_name, int) else sheet_name
        if aba not in livro.sheet_names:
            raise BipeInvalido('Bipe de produtos: escolha uma aba existente neste arquivo.')
        return _ler_aba(livro, aba)


def encontrar_aba_bipe(caminho, aba_escolhida=''):
    """Confere o conteúdo integral antes de aceitar o arquivo no envio guiado."""
    with pd.ExcelFile(caminho) as livro:
        if aba_escolhida and aba_escolhida not in livro.sheet_names:
            raise BipeInvalido('Escolha uma aba existente neste arquivo.')
        candidatas = []
        falhas = []
        for aba in ([aba_escolhida] if aba_escolhida else livro.sheet_names):
            try:
                _ler_aba(livro, aba)
                candidatas.append(aba)
            except BipeInvalido as exc:
                falhas.append(exc)
        if len(candidatas) == 1:
            return candidatas[0]
        if len(candidatas) > 1:
            raise BipeInvalido('Há mais de uma aba compatível com Bipe de produtos. Escolha a aba que deseja usar.')
        if falhas:
            raise max(falhas, key=lambda erro: erro.reconhecidos)
        raise BipeInvalido('Bipe de produtos: não foram encontradas abas com pedidos.')
