import io
import tempfile
import unittest
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from app import create_app
from app.cruzamento.engine import merge_tables, normalize, read_table

class EngineTests(unittest.TestCase):
    def test_chain_preserves_duplicates_unmatched_and_keys(self):
        key='0'+'1234567890'*4+'123'
        base=pd.DataFrame({'Pedido':['="1180000001"','1180000001','9','']})
        truck=pd.DataFrame({'Pedido':['1180000001.0',''], 'Chave':[key,'ERRADO']})
        a, stats=merge_tables(base,truck,['Pedido'],['Pedido'],['Chave'],'Carreta')
        self.assertEqual(len(a),4)
        self.assertEqual(stats['encontradas'],2)
        b,_=merge_tables(a,pd.DataFrame({'NF-e':[key],'Ocorrência':['Entregue']}),['Carreta · Chave'],['NF-e'],['Ocorrência'],'ESL')
        self.assertEqual(b['ESL · Ocorrência'].tolist(),['Entregue','Entregue','',''])
        self.assertEqual(b['Carreta · Chave'][0],key)

    def test_duplicate_policies_and_composite(self):
        base=pd.DataFrame({'NF':['1','1'],'CD':['590','1100']})
        source=pd.DataFrame({'Nota':['1','1','1'],'Centro':['590','590','1100'],'Status':['A','B','C']})
        with self.assertRaises(ValueError):
            merge_tables(base,source,['NF','CD'],['Nota','Centro'],['Status'],'Bipe')
        for policy, expected in [('first',['A','C']),('last',['B','C']),('join',['A | B','C']),('expand',['A','B','C'])]:
            result,_=merge_tables(base,source,['NF','CD'],['Nota','Centro'],['Status'],'Bipe',policy)
            self.assertEqual(result['Bipe · Status'].tolist(),expected)

    def test_csv_header_and_normalization(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'aleatorio.csv'
            p.write_text('Relatório;\nPedido;Situação\n00123;No depósito\n',encoding='utf-8-sig')
            df=read_table(p,header=2)
            self.assertEqual(df['Pedido'][0],'00123')
        self.assertEqual(normalize(' ="00123" '),'00123')

class WebTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.app=create_app()
        self.app.config.update(TESTING=True, SECRET_KEY='test-only', CRUZAMENTO_TEMP_DIR=self.temp.name)
        self.client=self.app.test_client()
        self.client.get('/cruzar-planilhas/')
    def tearDown(self): self.temp.cleanup()
    def post(self, data):
        with self.client.session_transaction() as session: data['csrf']=session['cruzamento_csrf']
        response=self.client.post('/cruzar-planilhas/',data=data,follow_redirects=True)
        self.assertEqual(response.status_code,200)
        return response
    def upload(self, df):
        buf=io.BytesIO()
        df.to_excel(buf,index=False,sheet_name='Dados');buf.seek(0)
        self.post({'action':'upload','file':(buf,'nome-qualquer.xlsx')})
        self.post({'action':'preview','sheet':'Dados','header':'1'})
    def test_full_flow_export_and_isolation(self):
        key='12345678901234567890123456789012345678901234'
        self.upload(pd.DataFrame({'Pedido':['1180000001','1180000001','9']}))
        self.post({'action':'base'})
        self.upload(pd.DataFrame({'Código':['1180000001'],'Chave':[key]}))
        self.post({'action':'merge','left0':'Pedido','right0':'Código','columns':['Chave'],'label':'Carreta','policy':'error'})
        self.upload(pd.DataFrame({'NF-e':[key],'Ocorrência':['=1+1']}))
        self.post({'action':'merge','left0':'Carreta · Chave','right0':'NF-e','columns':['Ocorrência'],'label':'ESL','policy':'error'})
        response=self.client.get('/cruzar-planilhas/baixar?numeric=Pedido')
        self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data));ws=book['Resultado']
        self.assertEqual(ws.max_row,4)
        self.assertEqual(ws['A2'].data_type,'n')
        self.assertEqual(ws['B2'].value,key)
        self.assertEqual(ws['C2'].data_type,'s')
        stranger=self.app.test_client()
        self.assertEqual(stranger.get('/cruzar-planilhas/baixar').status_code,302)
        self.post({'action':'finish'})
        self.assertEqual(self.client.get('/cruzar-planilhas/baixar').status_code,302)
        self.assertFalse(list(Path(self.temp.name).rglob('result.json')))
    def test_csrf_and_upload_validation(self):
        self.assertEqual(self.client.post('/cruzar-planilhas/',data={'action':'finish'}).status_code,400)
        response=self.post({'action':'upload','file':(io.BytesIO(b'hi'),'arquivo.exe')})
        self.assertIn('Selecione um arquivo',response.text)
    def test_routes_existing_modules(self):
        for url in ['/','/performance','/pagina_analise','/consulta_pedido']:
            self.assertEqual(self.client.get(url).status_code,200)

if __name__=='__main__': unittest.main()
